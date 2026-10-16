"""The blank workbook is empty; bench and discipline are one thing; the tagline is one text.

Covers the October 2026 changes to the workbooks:

* the blank workbook carries no example data on any sheet, including Shifts, Shift
  Requirements, Benches and Leave Types
* "Bench" and "Discipline" are interchangeable wherever a workbook mentions a section of the
  laboratory, and the examples use the eight benches the laboratory actually has
* the exported rota says "free" only when it was made on the free plan
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from labroster import api
from labroster.workbook import ERROR, read_workbook


def _modified(data: bytes, change) -> bytes:
    workbook = load_workbook(io.BytesIO(data))
    change(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _data_rows(workbook, name, header_row=2):
    return [[cell.value for cell in row if cell.value not in (None, "")]
            for row in workbook[name].iter_rows(min_row=header_row + 1)
            if any(cell.value not in (None, "") for cell in row)]


# --- the blank workbook is blank -----------------------------------------------------

@pytest.mark.parametrize("sheet", ["Staff", "Competencies", "Week Patterns", "Leave",
                                   "Shifts", "Shift Requirements", "Benches",
                                   "Leave Types"])
def test_every_data_sheet_of_the_blank_workbook_is_empty(sheet):
    workbook = load_workbook(io.BytesIO(api.blank_template_bytes()))
    assert _data_rows(workbook, sheet) == [], f"{sheet} should have no example data"


def test_the_blank_workbook_names_no_rota_or_laboratory():
    workbook = load_workbook(io.BytesIO(api.blank_template_bytes()))
    values = {str(row[0].value): row[1].value
              for row in workbook["Roster Details"].iter_rows(min_row=2)
              if row[0].value}
    for key in ("Rota name", "Organisation", "Department", "Site", "Prepared by"):
        assert values.get(key) in (None, ""), f"{key} should be blank, not {values[key]!r}"


def test_the_blank_workbook_still_has_its_headings_notes_and_dropdowns():
    workbook = load_workbook(io.BytesIO(api.blank_template_bytes()))
    assert [cell.value for cell in workbook["Shifts"][2]][:4] == [
        "Code", "Name", "Start", "End"]
    assert [cell.value for cell in workbook["Benches"][2]][:2] == ["Bench", "Discipline"]
    assert workbook["Shifts"].data_validations.dataValidation, "dropdowns were lost"
    assert "interchangeable" in str(workbook["Benches"]["A1"].value).lower()


def test_an_unfilled_blank_workbook_says_what_is_missing_rather_than_guessing():
    result = api.generate(api.blank_template_bytes())
    assert result["ok"] is False
    messages = " ".join(problem["message"] for problem in result["problems"])
    assert "No shifts are defined" in messages


def test_the_standard_absence_types_still_work_with_an_empty_leave_types_sheet():
    config, _ = read_workbook(io.BytesIO(api.blank_template_bytes()))
    assert {"al", "sl", "cl", "ml", "sd"} <= set(config.leave_types)


# --- bench and discipline are one thing --------------------------------------------

BENCHES = [("GRP", "Grouping"), ("XMTCH", "Crossmatch"),
           ("SPECS", "Specials Coagulation"), ("FBC", "Full Blood Count"),
           ("COAG", "Coagulation"), ("MORPH", "Morphology"),
           ("FLOW", "Flow Cytometry"), ("HB'THY", "Haemoglobinopathy")]


@pytest.mark.parametrize("builder", ["balanced_workbook_bytes", "demo_workbook_bytes"])
def test_the_examples_use_the_eight_named_benches(builder):
    workbook = load_workbook(io.BytesIO(getattr(api, builder)()))
    rows = [(row[0].value, row[1].value) for row in workbook["Benches"].iter_rows(min_row=3)
            if row[0].value]
    assert rows == BENCHES


@pytest.mark.parametrize("builder", ["blank_template_bytes", "balanced_workbook_bytes",
                                     "demo_workbook_bytes"])
def test_the_workbooks_say_bench_and_discipline_mean_the_same_thing(builder):
    workbook = load_workbook(io.BytesIO(getattr(api, builder)()))
    text = " ".join(str(cell.value) for name in ("Instructions", "Benches", "Competencies")
                    for row in workbook[name].iter_rows() for cell in row
                    if cell.value is not None).lower()
    assert "bench or discipline" in text
    assert "interchangeable" in text
    assert workbook["Competencies"][2][2].value == "Bench or Discipline"


def test_a_competency_may_name_the_bench_by_either_name():
    """GRP and Grouping are the same bench, so cover is identical whichever is written."""
    original = api.balanced_workbook_bytes()

    def spell_out(workbook):
        sheet = workbook["Competencies"]
        names = dict(BENCHES)
        for row in sheet.iter_rows(min_row=3):
            if row[2].value in names:
                row[2].value = names[row[2].value]

    spelled = _modified(original, spell_out)
    a, b = api.generate(original), api.generate(spelled)
    assert b["ok"], b.get("fatal")
    assert b["dashboard"]["critical_count"] == a["dashboard"]["critical_count"] == 0
    assert b["dashboard"]["review_count"] == a["dashboard"]["review_count"] == 0
    assert b["dashboard"]["shifts_meeting_all_requirements_percent"] == 100.0


def test_a_requirement_may_name_the_bench_by_either_name():
    def spell_out(workbook):
        sheet = workbook["Shift Requirements"]
        for row in sheet.iter_rows(min_row=3):
            for cell in row:
                if isinstance(cell.value, str) and "GRP:1" in cell.value:
                    cell.value = cell.value.replace("GRP:1", "Grouping:1")

    result = api.generate(_modified(api.balanced_workbook_bytes(), spell_out))
    assert result["ok"]
    assert result["dashboard"]["critical_count"] == 0


@pytest.mark.parametrize("heading", ["Discipline", "Bench", "Bench or Discipline"])
def test_the_competency_column_may_be_headed_any_of_the_three_ways(heading):
    def rename(workbook):
        workbook["Competencies"].cell(row=2, column=3, value=heading)

    config, problems = read_workbook(io.BytesIO(
        _modified(api.balanced_workbook_bytes(), rename)))
    assert [p for p in problems if p.severity == ERROR] == []
    assert config.competencies, "no competencies were read"


def test_a_workbook_using_one_word_for_both_is_unaffected():
    """Most laboratories write BT for the bench and BT for the discipline."""
    def one_word(workbook):
        sheet = workbook["Benches"]
        sheet["A3"], sheet["B3"] = "BT", "BT"
        comps = workbook["Competencies"]
        for row in comps.iter_rows(min_row=3):
            if row[2].value == "GRP":
                row[2].value = "BT"

    config, problems = read_workbook(io.BytesIO(
        _modified(api.balanced_workbook_bytes(), one_word)))
    assert [p for p in problems if p.severity == ERROR] == []
    assert any(bench.name == "BT" and bench.discipline == "BT" for bench in config.benches)


def test_the_rota_and_reports_use_the_bench_name():
    result = api.generate(api.balanced_workbook_bytes())
    manager = api.export_workbook(result["_scheduler"], result["_analysis"], "manager")
    workbook = load_workbook(io.BytesIO(manager))
    text = " ".join(str(cell.value) for sheet in ("Roster", "Competencies",
                                                  "Bench Allocations", "Competency Expiry")
                    for row in workbook[sheet].iter_rows() for cell in row
                    if cell.value is not None)
    assert "BENCH ALLOCATION" in text
    assert "SECTION ALLOCATION" not in text
    assert [c.value for c in workbook["Competencies"][3]][:2] == ["Name", "Bench"]
    assert "GROUPING" not in text, "a bench was shown by its upper-cased discipline"


# --- the tagline ----------------------------------------------------------------------

def test_the_product_tagline_is_the_agreed_text():
    from labroster import TAGLINE
    assert TAGLINE == ("A secure and intelligent workforce planning tool for diagnostic "
                       "laboratories.")


def _notes_text(data: bytes) -> str:
    workbook = load_workbook(io.BytesIO(data))
    return " ".join(str(cell.value) for row in workbook["Notes"].iter_rows()
                    for cell in row if cell.value is not None)


def test_a_rota_made_on_a_paid_plan_does_not_call_the_tool_free():
    result = api.generate(api.balanced_workbook_bytes())
    paid = _notes_text(api.export_workbook(
        result["_scheduler"], result["_analysis"], "staff", paid=True))
    assert "A secure and intelligent workforce planning tool" in paid
    assert "free" not in paid.lower()


def test_a_rota_made_on_the_free_plan_still_says_it_is_free():
    result = api.generate(api.balanced_workbook_bytes())
    free = _notes_text(api.export_workbook(result["_scheduler"], result["_analysis"], "staff"))
    assert "A free, secure and intelligent workforce planning tool" in free


def test_the_manager_report_follows_the_same_rule():
    result = api.generate(api.balanced_workbook_bytes())

    def instructions(paid):
        data = api.export_workbook(result["_scheduler"], result["_analysis"],
                                   "manager", paid=paid)
        sheet = load_workbook(io.BytesIO(data))["Instructions"]
        return " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)

    assert "free," not in instructions(True).lower()
    assert "A free, secure" in instructions(False)
