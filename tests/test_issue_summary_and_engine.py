"""A short list a manager can act on, and a roster the engine does not leave half done.

The first half covers how findings are gathered into one item per cause, and two advisory
checks that were saying more than they knew. The second covers the scheduler: bench cover
found by matching, a same-day repair, and sharing hours by share of contract.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from conftest import (
    MONDAY, competent, make_config, make_shift, make_staff,
)
from labroster import api
from labroster.analysis import CRITICAL, PASSED, REVIEW, Analysis
from labroster.models import Bench, Rules, ShiftRequirement
from labroster.scheduler import Scheduler


def build(config):
    scheduler = Scheduler(config)
    scheduler.build()
    return scheduler, Analysis(scheduler)


# --- one item per cause -----------------------------------------------------------------

def _short_staffed(days):
    return make_config([make_staff("A")],
                       requirements=[ShiftRequirement("C", min_staff=2)], days=days)


def test_a_cause_repeated_on_many_dates_is_one_item():
    _, analysis = build(_short_staffed(days=9))
    raw = [i for i in analysis.issues if i.severity == CRITICAL]
    summary = [i for i in analysis.summary_issues if i.severity == CRITICAL]
    assert len(raw) == 9                       # nothing is lost from the record
    assert len(summary) == 1                   # but the manager sees one thing to fix
    item = summary[0]
    assert item.count == 9
    assert "9 shifts" in item.title
    assert item.details and "9 dates" in item.details[0]["explanation"]
    assert analysis.metrics["critical_count"] == 1
    assert analysis.metrics["critical_detail_count"] == 9


def test_a_single_finding_is_shown_exactly_as_it_was():
    _, analysis = build(_short_staffed(days=1))
    summary = [i for i in analysis.summary_issues if i.severity == CRITICAL]
    assert len(summary) == 1
    assert summary[0].count == 1 and summary[0].details == []
    assert summary[0].title == "C is short of staff"


def test_the_title_counts_shifts_not_findings():
    """One shift can fail two ways at once; it is still one shift."""
    staff = [make_staff("A"), make_staff("B")]
    config = make_config(
        staff, requirements=[ShiftRequirement("C", min_staff=2, min_senior=1)], days=4)
    _, analysis = build(config)
    for item in analysis.summary_issues:
        if item.severity == CRITICAL and item.details:
            assert f"{item.count} shifts" not in item.title or item.count == 4


def test_different_causes_stay_separate_items():
    """Core lacks a senior every day; Early is short of staff every day."""
    shifts = [make_shift("C"), make_shift("E", start="07:00", end="15:00")]
    config = make_config(
        [make_staff(sid) for sid in "ABC"], shifts=shifts,
        requirements=[ShiftRequirement("C", min_staff=1, min_senior=1),
                      ShiftRequirement("E", min_staff=5)], days=5)
    _, analysis = build(config)
    items = {i.category: i for i in analysis.summary_issues if i.severity == CRITICAL}
    assert {"Staffing level", "Senior cover"} <= set(items)
    assert items["Senior cover"].count == items["Staffing level"].count == 5


def test_passed_checks_are_never_merged():
    result = api.generate(api.balanced_workbook_bytes())
    passed = [i for i in result["issues"] if i["severity"] == "PASSED"]
    assert len(passed) == result["dashboard"]["passed_count"] >= 5
    assert all(i["details"] == [] for i in passed)


def test_the_exported_issues_sheet_lists_the_findings_under_each_item():
    result = api.generate(api.demo_workbook_bytes())
    data = api.export_workbook(result["_scheduler"], result["_analysis"], "manager")
    sheet = load_workbook(io.BytesIO(data))["Issues"]
    lines = [str(row[2].value) for row in sheet.iter_rows() if row[2].value]
    assert any(line.lstrip().startswith("•") for line in lines), \
        "the findings behind a consolidated item should be listed beneath it"
    headers = next([c.value for c in row] for row in sheet.iter_rows()
                   if row[0].value == "Severity")
    assert "Bench" in headers and "Section" not in headers


# --- advisory checks that said more than they knew -----------------------------------

def _long_period_config(days):
    shift = make_shift("C")
    staff = [make_staff("A"), make_staff("B")]
    comps = [competent("A", "X"), competent("A", "Y"), competent("B", "X")]
    return make_config(
        staff, shifts=[shift],
        requirements=[ShiftRequirement("C", min_staff=2)],
        benches=[Bench(name="X", discipline="X")], competencies=comps, days=days,
        rules=Rules(rotation_warning_days=56))


def test_rotation_is_not_claimed_over_a_roster_shorter_than_the_interval():
    """Not having worked a bench in 31 days does not breach a 56-day rotation."""
    _, analysis = build(_long_period_config(days=31))
    assert [i for i in analysis.issues if i.category == "Section rotation"] == []


def test_rotation_is_reported_once_the_roster_is_long_enough_to_judge_it():
    _, analysis = build(_long_period_config(days=60))
    flagged = [i for i in analysis.issues if i.category == "Section rotation"]
    assert flagged and "Y" in flagged[0].title


def _resilience_config(with_backup, min_staff=2):
    shifts = [make_shift("C"), make_shift("E", start="07:00", end="15:00")]
    staff = [make_staff("A"), make_staff("B"), make_staff("D")]
    comps = [competent("A", "X")] + ([competent("D", "X")] if with_backup else [])
    requirements = [
        ShiftRequirement("C", min_staff=min_staff, required_competencies={"X": 1}),
        ShiftRequirement("E", min_staff=1, required_competencies=(
            {"X": 1} if with_backup else {})),
    ]
    return make_config(staff, shifts=shifts, requirements=requirements,
                       competencies=comps, days=1)


def test_one_competent_person_with_nobody_to_stand_in_is_flagged():
    _, analysis = build(_resilience_config(with_backup=False))
    assert [i for i in analysis.issues if i.category == "Shift resilience"]


def test_one_competent_person_is_not_flagged_when_somebody_else_on_the_roster_could_stand_in():
    _, analysis = build(_resilience_config(with_backup=True))
    assert [i for i in analysis.issues if i.category == "Shift resilience"] == []


def test_a_single_handed_shift_is_not_reported_as_fragile():
    """The manager chose to run it with one person; saying so back tells them nothing."""
    _, analysis = build(_resilience_config(with_backup=False, min_staff=1))
    assert [i for i in analysis.issues if i.category == "Shift resilience"] == []


# --- the scheduler -----------------------------------------------------------------------

def test_benches_are_covered_whenever_the_people_on_duty_allow_it():
    """Handing benches out one at a time strands one: P:[A,C], Q:[A,B], R:[B].

    Taking the bench with fewest people first gives R to B, then P to A, leaving Q with
    nobody. A full assignment exists (C-P, A-Q, B-R), and matching finds it.
    """
    staff = [make_staff(sid) for sid in "ABC"]
    comps = [competent("A", "P"), competent("A", "Q"), competent("B", "Q"),
             competent("B", "R"), competent("C", "P")]
    benches = [Bench(name=name, discipline=name) for name in "PQR"]
    config = make_config(staff, shifts=[make_shift("C")],
                         requirements=[ShiftRequirement("C", min_staff=3)],
                         benches=benches, competencies=comps, days=1)
    scheduler, analysis = build(config)
    assert [s for s in scheduler.shortfalls if s.kind == "bench"] == []
    at = {allocation.bench_name: allocation.staff_id
          for allocation in scheduler.bench_allocations}
    assert at == {"R": "B", "Q": "A", "P": "C"}
    assert not [i for i in analysis.issues
                if i.severity == CRITICAL and "coverage" in i.category.lower()]


@pytest.mark.parametrize("seed", range(12))
def test_the_one_person_a_shift_needs_is_not_used_up_by_an_earlier_shift(seed):
    """Early and Core both want a free person; only A can stand at the X bench on Core.

    Whichever order the shifts happen to be filled in, A must end up on Core.
    """
    shifts = [make_shift("E", start="07:00", end="15:00"), make_shift("C")]
    config = make_config(
        [make_staff("A"), make_staff("B")], shifts=shifts,
        requirements=[ShiftRequirement("E", min_staff=1),
                      ShiftRequirement("C", min_staff=1)],
        benches=[Bench(name="X", discipline="X", shift_codes=["C"])],
        competencies=[competent("A", "X")], days=1, rules=Rules(seed=seed))
    scheduler, analysis = build(config)
    assert scheduler.assignments[(MONDAY, "A")].shift_code == "C"
    assert scheduler.assignments[(MONDAY, "B")].shift_code == "E"
    assert [s for s in scheduler.shortfalls if s.kind in ("bench", "competency")] == []


def test_part_timers_and_full_timers_end_up_the_same_share_of_their_contract():
    """Ranked by hours owed in absolute terms, full-timers always went first."""
    staff = [make_staff("F1"), make_staff("F2"),
             make_staff("P1", contracted_weekly_hours=22.5, fte=0.6),
             make_staff("P2", contracted_weekly_hours=22.5, fte=0.6)]
    # Two on every day is 56 person-days against 60 contracted: enough work, and few enough
    # people on each day that who works is a real choice.
    config = make_config(staff, days=28,
                         requirements=[ShiftRequirement("C", min_staff=2)])
    scheduler, analysis = build(config)
    shares = [scheduler.by_id[sid].total_accounted_hours
              / scheduler.by_id[sid].target_period_hours for sid in ("F1", "F2", "P1", "P2")]
    assert max(shares) - min(shares) <= 0.08, shares
    assert analysis.metrics["staff_outside_target_hours"] == 0


def test_balancing_never_breaks_cover_or_a_managers_own_placement():
    from labroster.scheduler import Assignment
    staff = [make_staff("A"), make_staff("B"), make_staff("C")]
    manual = [Assignment(day=MONDAY, staff_id="A", shift_code="C")]
    config = make_config(staff, requirements=[ShiftRequirement("C", min_staff=2)],
                         days=14)
    scheduler = Scheduler(config, manual_assignments=manual)
    scheduler.build()
    assert scheduler.assignments[(MONDAY, "A")].is_manual
    for day in scheduler.days:
        assert len(scheduler.assigned_to(day, config.shifts[0])) >= 2
